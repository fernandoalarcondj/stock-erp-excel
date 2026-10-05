"""Lectura de los PDF del ERP (Stock Puro / Stock Destinado a Clientes),
cálculo de KPIs y generación de Excel y CSV.

Principio de diseño: una fila nunca se descarta por tener un dato desconocido.
Producto, calidad, alistamiento u observación que no se reconozcan quedan en
blanco y el texto sobrante se guarda en la columna "Texto sin reconocer".
"""
import io
import re

import pandas as pd
import pdfplumber
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from catalogos import (ALISTAMIENTOS, CALIDADES, CARTULINA_GENERICA,
                       CARTULINA_SEGUN_CALIDAD, OBSERVACIONES, PRODUCTOS)

NUM = r"\d+(?:\.\d+)?"
FUENTE = "Arial"
SIN_DATO = "(sin dato)"


# ----------------------------------------------------------------------------
# Catálogos -> patrones de búsqueda
# ----------------------------------------------------------------------------
def _patron(alias):
    partes = r"\s+".join(re.escape(p) for p in alias.split())
    return re.compile(r"(?<![A-Z0-9])" + partes + r"(?![A-Z0-9])")


def _compilar(catalogo):
    items = [(a, canon) for canon, als in catalogo.items() for a in als]
    items.sort(key=lambda x: -len(x[0]))  # primero los alias más largos
    return [(_patron(a), canon) for a, canon in items]


_RE_CAL = _compilar(CALIDADES)
_RE_PROD = _compilar(PRODUCTOS)
_RE_OBS = _compilar(OBSERVACIONES)
_RE_CART = _compilar({"x": CARTULINA_GENERICA})
_ALIST = {a: canon for canon, als in ALISTAMIENTOS.items() for a in als}


def _buscar(texto, compilados):
    for pat, canon in compilados:
        m = pat.search(texto)
        if m:
            return canon, texto[: m.start()] + " " + texto[m.end():]
    return "", texto


def interpretar_mid(mid):
    """Del texto entre alistamiento y números saca producto, calidad y obs."""
    u = re.sub(r"\s+", " ", (mid or "").upper()).strip()
    obs, u = _buscar(u, _RE_OBS)
    calidad, u = _buscar(u, _RE_CAL)
    producto, u = _buscar(u, _RE_PROD)
    if not producto:
        gen, u2 = _buscar(u, _RE_CART)
        if gen:
            producto, u = CARTULINA_SEGUN_CALIDAD.get(calidad, ""), u2
    return producto, calidad, obs, " ".join(u.split())


def _fecha_y_resto(head):
    """El lote/orden y la fecha pueden venir superpuestos y partidos
    ('087543/2 2 6/11/2025'). Sin espacios, la fecha son los últimos 10 chars."""
    s = head.replace(" ", "")
    fecha = pd.to_datetime(s[-10:], format="%d/%m/%Y", errors="coerce")
    if pd.isna(fecha):
        return pd.NaT, s
    return fecha, s[:-10]


# ----------------------------------------------------------------------------
# Lectura del PDF
# ----------------------------------------------------------------------------
def extraer_lineas(archivo):
    lineas = []
    with pdfplumber.open(archivo) as pdf:
        for pagina in pdf.pages:
            txt = pagina.extract_text() or ""
            lineas.extend(l.strip() for l in txt.splitlines() if l.strip())
    return lineas


def detectar_tipo(lineas):
    cab = " ".join(lineas[:15]).lower()
    if "stock puro" in cab:
        return "puro"
    if "destinado" in cab:
        return "destinado"
    return None


def _fecha_corte(lineas):
    for l in lineas[:15]:
        m = re.search(r"-\s*(\d{2}/\d{2}/\d{4})", l)
        if m:
            return pd.to_datetime(m.group(1), format="%d/%m/%Y").date()
    return None


def _num_ar(txt):
    return float(txt.replace(".", "").replace(",", "."))


RE_TIENE_FECHA = re.compile(r"\d{1,2}/\d{2}/\d{4}")

# ---- Destinado ---------------------------------------------------------------
RE_CLIENTE = re.compile(r"^Cliente:\s*(?P<nombre>.*?)\s*\((?P<cod>\d+)\)\s*(?P<dir>.*)$")
RE_TOT_CLIENTE = re.compile(
    r"^Total Kg Cliente:\s*(?P<kg>\d+)\s+Total Unidad Cliente:\s*(?P<u>\d+)"
    r"\s+Valorizaci\S*:\s*(?P<val>[\d.,]+)"
)
RE_FILA_D = re.compile(
    rf"^(?P<head>.+?)\s+(?P<alist>[A-Z]+)\s+(?:(?P<mid>.*?)\s+)?"
    rf"(?P<unid>\d+)\s+(?P<kilos>\d+)\s+(?P<gram>\d+)\s+(?P<corte>{NUM})\s+"
    rf"(?P<diam>\d+)\s+(?P<mon>\$|USD)\s*(?P<precio>{NUM})\s+(?P<pedido>\d{{6}}-\d{{2}})\s*$"
)
COLS_D = ["Cliente", "Cód. Cliente", "Dirección", "Orden Fabr.", "Fecha", "Alistamiento",
          "Producto", "Calidad", "Observ.", "Unidades", "Kilos", "Gramaje", "Corte",
          "Diámetro", "Moneda", "Precio", "Pedido", "Valorización", "Texto sin reconocer"]


def parsear_destinado(lineas):
    filas, totales, no_leidas = [], {}, []
    cliente = cod = direccion = None
    for l in lineas:
        m = RE_CLIENTE.match(l)
        if m:
            cliente, cod, direccion = m["nombre"], m["cod"], m["dir"]
            continue
        m = RE_TOT_CLIENTE.match(l)
        if m:
            totales[cod] = dict(cliente=cliente, kg=int(m["kg"]), unid=int(m["u"]))
            continue
        m = RE_FILA_D.match(l)
        if not m:
            if RE_TIENE_FECHA.search(l) and not l.startswith("Periodo"):
                no_leidas.append(l)
            continue
        fecha, orden = _fecha_y_resto(m["head"])
        prod, cal, obs, resto = interpretar_mid(m["mid"])
        filas.append({
            "Cliente": cliente, "Cód. Cliente": cod, "Dirección": direccion,
            "Orden Fabr.": orden, "Fecha": fecha,
            "Alistamiento": _ALIST.get(m["alist"], ""),
            "Producto": prod, "Calidad": cal, "Observ.": obs,
            "Unidades": int(m["unid"]), "Kilos": int(m["kilos"]),
            "Gramaje": int(m["gram"]), "Corte": float(m["corte"]),
            "Diámetro": int(m["diam"]),
            "Moneda": "USD" if m["mon"] == "USD" else "ARS",
            "Precio": float(m["precio"]), "Pedido": m["pedido"],
            "Texto sin reconocer": (resto + (" | " + m["alist"] if m["alist"] not in _ALIST else "")).strip(" |"),
        })
    df = pd.DataFrame(filas, columns=[c for c in COLS_D if c != "Valorización"])
    if not df.empty:
        df["Valorización"] = df["Kilos"] * df["Precio"]
        df = df[COLS_D]
    avisos = []
    if not df.empty:
        real = df.groupby("Cód. Cliente")[["Kilos", "Unidades"]].sum()
        for c, t in totales.items():
            if c not in real.index:
                avisos.append(f"{t['cliente']}: figura en el PDF pero no se leyeron filas.")
            elif real.loc[c, "Kilos"] != t["kg"] or real.loc[c, "Unidades"] != t["unid"]:
                avisos.append(f"{t['cliente']}: el PDF indica {t['kg']} kg / {t['unid']} u y se leyeron "
                              f"{real.loc[c, 'Kilos']} kg / {real.loc[c, 'Unidades']} u.")
    tot = {"kg": sum(t["kg"] for t in totales.values()), "unid": sum(t["unid"] for t in totales.values())}
    return df, avisos, tot, no_leidas


# ---- Puro --------------------------------------------------------------------
RE_FILA_P = re.compile(
    rf"^(?P<head>.+?)\s+(?P<alist>[A-Z]+)\s+(?:(?P<mid>.*?)\s+)?"
    rf"(?P<kilos>\d+)\s+(?P<unid>\d+)\s+(?P<gram>\d+)\s+(?P<corte>{NUM})\s+(?P<diam>\d+)\s*$"
)
RE_TOT_CALIDAD = re.compile(r"^Total por Calidad:\s*(?P<kg>\d+)\s+(?P<u>\d+)")
COLS_P = ["Lote", "Sub-lote", "Fecha", "Alistamiento", "Producto", "Calidad", "Observ.",
          "Kilos", "Unidades", "Gramaje", "Corte", "Diámetro", "Texto sin reconocer"]


def parsear_puro(lineas):
    filas, no_leidas, tot_kg, tot_u = [], [], 0, 0
    for l in lineas:
        m = RE_TOT_CALIDAD.match(l)
        if m:
            tot_kg += int(m["kg"])
            tot_u += int(m["u"])
            continue
        m = RE_FILA_P.match(l)
        if not m:
            if RE_TIENE_FECHA.search(l) and not l.startswith("Periodo"):
                no_leidas.append(l)
            continue
        fecha, resto_head = _fecha_y_resto(m["head"])
        lote, _, sub = resto_head.partition("/")
        prod, cal, obs, resto = interpretar_mid(m["mid"])
        filas.append({
            "Lote": lote, "Sub-lote": sub, "Fecha": fecha,
            "Alistamiento": _ALIST.get(m["alist"], ""),
            "Producto": prod, "Calidad": cal, "Observ.": obs,
            "Kilos": int(m["kilos"]), "Unidades": int(m["unid"]),
            "Gramaje": int(m["gram"]), "Corte": float(m["corte"]),
            "Diámetro": int(m["diam"]),
            "Texto sin reconocer": (resto + (" | " + m["alist"] if m["alist"] not in _ALIST else "")).strip(" |"),
        })
    df = pd.DataFrame(filas, columns=COLS_P)
    avisos = []
    if not df.empty and (df["Kilos"].sum() != tot_kg or df["Unidades"].sum() != tot_u):
        avisos.append(f"Los totales del PDF son {tot_kg} kg / {tot_u} u, pero se leyeron "
                      f"{df['Kilos'].sum()} kg / {df['Unidades'].sum()} u.")
    return df, avisos, {"kg": tot_kg, "unid": tot_u}, no_leidas


def procesar_pdf(archivo):
    """Devuelve dict: tipo, df, avisos, totales_pdf, fecha_corte."""
    lineas = extraer_lineas(archivo)
    tipo = detectar_tipo(lineas)
    if tipo == "destinado":
        df, avisos, tot, no_leidas = parsear_destinado(lineas)
    elif tipo == "puro":
        df, avisos, tot, no_leidas = parsear_puro(lineas)
    else:
        raise ValueError("No se reconoce el PDF: no es un reporte de Stock Puro ni de Stock Destinado a Clientes.")
    if df.empty:
        raise ValueError("No se encontraron filas de datos en el PDF.")
    for col in ("Producto", "Calidad", "Alistamiento"):
        n = int((df[col] == "").sum())
        if n:
            avisos.append(f"{n} fila(s) con «{col}» no reconocido: quedó en blanco (ver columna «Texto sin reconocer»).")
    if no_leidas:
        avisos.append(f"{len(no_leidas)} línea(s) con fecha no se pudieron interpretar y no están en el resultado: "
                      + " || ".join(no_leidas[:5]))
    return {"tipo": tipo, "df": df, "avisos": avisos, "totales_pdf": tot, "fecha_corte": _fecha_corte(lineas)}


# ----------------------------------------------------------------------------
# KPIs (los mismos que la hoja KPIs del Excel, calculados en pandas para la app)
# ----------------------------------------------------------------------------
def calcular_kpis(df, fecha_corte=None, dias_antiguo=90):
    k, u = df["Kilos"].sum(), df["Unidades"].sum()
    fc = pd.Timestamp(fecha_corte) if fecha_corte else df["Fecha"].max()
    antig = (fc - df["Fecha"]).dt.days
    viejo = df.loc[antig > dias_antiguo, "Kilos"].sum()
    out = {
        "Registros": len(df),
        "Unidades": int(u),
        "Kilos": int(k),
        "Kilos por unidad": k / u if u else 0,
        "Gramaje prom. ponderado (g/m²)": (df["Kilos"] * df["Gramaje"]).sum() / k if k else 0,
        "Corte prom. ponderado (cm)": (df["Kilos"] * df["Corte"]).sum() / k if k else 0,
        f"Kilos con más de {dias_antiguo} días": int(viejo),
        f"% con más de {dias_antiguo} días": viejo / k if k else 0,
        "Antigüedad prom. ponderada (días)": (antig * df["Kilos"]).sum() / k if k else 0,
    }
    if "Cliente" in df:
        out["Clientes"] = df["Cliente"].nunique()
        v = df.groupby("Moneda")["Valorización"].sum()
        out["Valorización ARS"] = v.get("ARS", 0.0)
        out["Valorización USD"] = v.get("USD", 0.0)
    return out


# ----------------------------------------------------------------------------
# CSV
# ----------------------------------------------------------------------------
def generar_csv(df):
    """CSV pensado para Excel en español: separador ';', decimales con coma."""
    d = df.copy()
    d["Fecha"] = d["Fecha"].dt.strftime("%d/%m/%Y")
    return d.to_csv(index=False, sep=";", decimal=",", encoding="utf-8-sig").encode("utf-8-sig")


# ----------------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUB_FILL = PatternFill("solid", fgColor="DDEBF7")
F_NORMAL = Font(name=FUENTE, size=10)
F_BOLD = Font(name=FUENTE, size=10, bold=True)
F_INPUT = Font(name=FUENTE, size=10, color="0000FF")
F_HDR = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")


def _encabezado(ws, fila, ncols, col0=1):
    for c in range(col0, col0 + ncols):
        cel = ws.cell(row=fila, column=c)
        cel.font, cel.fill = F_HDR, HEADER_FILL
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _formatear_tabla(ws):
    _encabezado(ws, 1, ws.max_column)
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.font = F_NORMAL
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.columns, 1):
        mx = max((len(str(c.value)) for c in col
                  if c.value is not None and not str(c.value).startswith("=")), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(mx + 2, 11), 45)


def _hoja_resumen(w, df_res, nombre, cols_suma):
    """Resumen con TOTAL por fórmula (si hay Moneda, totales separados ARS/USD)."""
    df_res.to_excel(w, sheet_name=nombre, index=False)
    ws = w.sheets[nombre]
    cols, n = list(df_res.columns), len(df_res) + 1
    L = lambda c: get_column_letter(cols.index(c) + 1)
    filas_tot = []
    if "Moneda" in cols:
        rm = f"${L('Moneda')}$2:${L('Moneda')}${n}"
        for mon in ("ARS", "USD"):
            r = n + 1 + len(filas_tot)
            ws.cell(row=r, column=1, value=f"TOTAL {mon}")
            for c in cols_suma:
                ws.cell(row=r, column=cols.index(c) + 1,
                        value=f'=SUMIF({rm},"{mon}",{L(c)}2:{L(c)}{n})')
            filas_tot.append(r)
    else:
        r = n + 1
        ws.cell(row=r, column=1, value="TOTAL")
        for c in cols_suma:
            ws.cell(row=r, column=cols.index(c) + 1, value=f"=SUM({L(c)}2:{L(c)}{n})")
        filas_tot.append(r)
    _formatear_tabla(ws)
    for r in filas_tot:
        for c in ws[r]:
            c.font = F_BOLD
    for c in cols_suma:
        for r in range(2, n + 1 + len(filas_tot)):
            ws.cell(row=r, column=cols.index(c) + 1).number_format = (
                "#,##0.00" if c == "Valorización" else "#,##0")


def _hoja_kpis(wb, df, tipo, fecha_corte):
    """Hoja KPIs: todo con fórmulas que apuntan a la hoja Detalle."""
    ws = wb.create_sheet("KPIs", 0)
    cols, n = list(df.columns), len(df) + 1
    rng = lambda c: f"Detalle!${get_column_letter(cols.index(c) + 1)}$2:${get_column_letter(cols.index(c) + 1)}${n}"
    K, U, G, C, F = rng("Kilos"), rng("Unidades"), rng("Gramaje"), rng("Corte"), rng("Fecha")

    ws["A1"] = ("KPIs – Stock Puro" if tipo == "puro" else "KPIs – Stock Destinado a Clientes")
    ws["A1"].font = Font(name=FUENTE, size=14, bold=True)
    ws["A3"], ws["B3"] = "Fecha de corte del reporte", fecha_corte
    ws["A4"], ws["B4"] = "Días para considerar stock antiguo", 90
    ws["C3"], ws["C4"] = "Dato tomado del PDF (editable)", "Parámetro editable"
    ws["B3"].number_format = "DD/MM/YYYY"
    for a in ("B3", "B4"):
        ws[a].font = F_INPUT
        ws[a].fill = PatternFill("solid", fgColor="FFFF00")
    for a in ("A3", "A4", "C3", "C4"):
        ws[a].font = F_NORMAL

    ws["A6"], ws["B6"], ws["C6"] = "Indicador", "Valor", "Detalle"
    _encabezado(ws, 6, 3)
    k = [
        ("Registros (filas)", f"=COUNT({K})", "#,##0", "Cantidad de líneas del reporte"),
        ("Unidades totales", f"=SUM({U})", "#,##0", ""),
        ("Kilos totales", f"=SUM({K})", "#,##0", ""),
        ("Kilos por unidad", "=IFERROR(B9/B8,0)", "#,##0.0", "Peso promedio por unidad"),
        ("Gramaje promedio ponderado", f"=IFERROR(SUMPRODUCT({K},{G})/B9,0)", "#,##0.0", "Ponderado por kilos"),
        ("Corte promedio ponderado", f"=IFERROR(SUMPRODUCT({K},{C})/B9,0)", "#,##0.0", "Ponderado por kilos (cm)"),
        ("Kilos con antigüedad mayor al parámetro", f'=SUMIF({F},"<"&($B$3-$B$4),{K})', "#,##0", "Según B3 y B4"),
        ("% de kilos antiguos", "=IFERROR(B13/B9,0)", "0.0%", ""),
        ("Antigüedad promedio ponderada (días)", f"=IFERROR($B$3-SUMPRODUCT({K},{F})/B9,0)", "#,##0", "Ponderada por kilos"),
        ("Filas con Producto en blanco", f"=COUNTBLANK({rng('Producto')})", "#,##0", "Control de calidad de datos"),
        ("Filas con Calidad en blanco", f"=COUNTBLANK({rng('Calidad')})", "#,##0", "Control de calidad de datos"),
        ("Filas con Alistamiento en blanco", f"=COUNTBLANK({rng('Alistamiento')})", "#,##0", "Control de calidad de datos"),
    ]
    if tipo == "destinado":
        cl, mo, va = rng("Cliente"), rng("Moneda"), rng("Valorización")
        k += [
            ("Clientes distintos", f'=SUMPRODUCT(({cl}<>"")/COUNTIF({cl},{cl}&""))', "#,##0", ""),
            ("Valorización ARS", f'=SUMIF({mo},"ARS",{va})', "#,##0.00", "Kilos × precio"),
            ("Valorización USD", f'=SUMIF({mo},"USD",{va})', "#,##0.00", "Kilos × precio (no se mezcla con ARS)"),
        ]
    for i, (lab, f, fmt, nota) in enumerate(k):
        r = 7 + i
        ws.cell(row=r, column=1, value=lab).font = F_NORMAL
        c = ws.cell(row=r, column=2, value=f)
        c.font, c.number_format = F_BOLD, fmt
        ws.cell(row=r, column=3, value=nota).font = F_NORMAL

    # Tablas de desglose (SUMIF contra Detalle). "(sin dato)" = total - resto.
    def tabla(r, titulo, col, valores, hay_blancos, total_txt="TOTAL"):
        ws.cell(row=r, column=1, value=titulo).font = Font(name=FUENTE, size=11, bold=True)
        for j, h in enumerate([col, "Unidades", "Kilos", "% de kilos"], 1):
            ws.cell(row=r + 1, column=j, value=h)
        _encabezado(ws, r + 1, 4)
        crit = rng(col)
        r0 = r + 2
        for i, v in enumerate(valores):
            rr = r0 + i
            ws.cell(row=rr, column=1, value=v)
            ws.cell(row=rr, column=2, value=f"=SUMIF({crit},A{rr},{U})")
            ws.cell(row=rr, column=3, value=f"=SUMIF({crit},A{rr},{K})")
        rr = r0 + len(valores)
        ult = rr - 1
        if hay_blancos:
            ws.cell(row=rr, column=1, value=SIN_DATO)
            ws.cell(row=rr, column=2, value=f"=$B$8-SUM(B{r0}:B{ult})" if valores else "=$B$8")
            ws.cell(row=rr, column=3, value=f"=$B$9-SUM(C{r0}:C{ult})" if valores else "=$B$9")
            rr += 1
        ws.cell(row=rr, column=1, value=total_txt)
        ws.cell(row=rr, column=2, value=f"=SUM(B{r0}:B{rr - 1})")
        ws.cell(row=rr, column=3, value=f"=SUM(C{r0}:C{rr - 1})")
        for x in range(r0, rr + 1):
            ws.cell(row=x, column=4, value=f"=IFERROR(C{x}/$B$9,0)")
            for cc, fmt in ((2, "#,##0"), (3, "#,##0"), (4, "0.0%")):
                ws.cell(row=x, column=cc).number_format = fmt
            for cc in range(1, 5):
                ws.cell(row=x, column=cc).font = F_BOLD if x == rr else F_NORMAL
        return rr + 3

    r = 7 + len(k) + 2
    for titulo, col in (("Kilos por Producto", "Producto"), ("Kilos por Calidad", "Calidad"),
                        ("Kilos por Alistamiento", "Alistamiento"), ("Kilos por Gramaje", "Gramaje")):
        s = df[col]
        vals = sorted(v for v in s.unique() if v != "" and not pd.isna(v))
        r = tabla(r, titulo, col, vals, bool((s == "").any()))
    if tipo == "destinado":
        top = df.groupby("Cliente")["Kilos"].sum().sort_values(ascending=False).head(10).index.tolist()
        tabla(r, "Top 10 clientes por kilos", "Cliente", top, False, "TOTAL TOP 10")

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 38
    ws.column_dimensions["D"].width = 14


def generar_excel(df, tipo, fecha_corte=None):
    d = df.copy()
    for c in d.select_dtypes(include="object").columns:
        d[c] = d[c].where(d[c].fillna("") != "", None)  # blancos reales en Excel
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="DD/MM/YYYY") as w:
        d.to_excel(w, sheet_name="Detalle", index=False)
        ws = w.sheets["Detalle"]
        cols = list(d.columns)
        _formatear_tabla(ws)
        ws.auto_filter.ref = ws.dimensions
        for col in ("Kilos", "Unidades"):
            i = cols.index(col) + 1
            for r in range(2, len(d) + 2):
                ws.cell(row=r, column=i).number_format = "#,##0"
        if tipo == "destinado":  # Valorización como fórmula Kilos × Precio
            ck = get_column_letter(cols.index("Kilos") + 1)
            cp = get_column_letter(cols.index("Precio") + 1)
            iv = cols.index("Valorización") + 1
            for r in range(2, len(d) + 2):
                c = ws.cell(row=r, column=iv, value=f"={ck}{r}*{cp}{r}")
                c.number_format, c.font = "#,##0.00", F_NORMAL
        g = df.copy()
        for c in ("Producto", "Calidad", "Alistamiento"):
            g[c] = g[c].replace("", SIN_DATO)
        ag = dict(Unidades=("Unidades", "sum"), Kilos=("Kilos", "sum"))
        if tipo == "destinado":
            ag["Valorización"] = ("Valorización", "sum")
            _hoja_resumen(w, g.groupby(["Cliente", "Cód. Cliente", "Moneda"], as_index=False).agg(**ag),
                          "Resumen Cliente", ["Unidades", "Kilos", "Valorización"])
            _hoja_resumen(w, g.groupby(["Producto", "Calidad", "Gramaje", "Moneda"], as_index=False).agg(**ag),
                          "Resumen Producto", ["Unidades", "Kilos", "Valorización"])
        else:
            _hoja_resumen(w, g.groupby(["Alistamiento", "Producto", "Calidad"], as_index=False).agg(**ag),
                          "Resumen Producto", ["Unidades", "Kilos"])
            _hoja_resumen(w, g.groupby(["Producto", "Calidad", "Gramaje", "Corte"], as_index=False).agg(**ag),
                          "Resumen Gramaje y Corte", ["Unidades", "Kilos"])
        fc = pd.Timestamp(fecha_corte) if fecha_corte else df["Fecha"].max()
        _hoja_kpis(w.book, df, tipo, fc.to_pydatetime())
    return buf.getvalue()
